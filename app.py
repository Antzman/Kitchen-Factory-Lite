from __future__ import annotations

import csv
import hmac
import io
import os
import secrets
import time
from datetime import datetime
from decimal import Decimal
from threading import Lock

from flask import Flask, flash, redirect, render_template, request, session, url_for

from db import (
    clear_tenant_context,
    get_db_connection,
    get_secret_key,
    init_db,
    seed_data,
    set_tenant_context,
)
from services import (
    add_audit,
    audit_entries,
    create_bulk_portioning,
    create_category,
    create_manufacturing_transaction,
    create_stock_item,
    create_yield_loss_portioning,
    dashboard_metrics,
    get_user_by_username,
    get_user_by_id,
    authenticate_user,
    register_company,
    change_temporary_password,
    get_settings,
    get_setting,
    import_stock_items,
    list_categories,
    list_manufacturing_history,
    list_recent_portioning,
    list_stock_items,
    stock_item_report,
    parse_csv,
    set_stock_item_active,
    set_setting,
    update_stock_item,
    validate_import_rows,
    update_settings,
    list_users,
    create_user,
    update_user,
    administrator_reset_user_password,
    emergency_admin_recovery,
    update_category,
    set_category_active,
    list_menu_categories,
    create_menu_category,
    update_menu_category,
    set_menu_category_active,
    delete_menu_category,
    list_menu_items,
    get_menu_item,
    get_menu_item_recipe,
    create_menu_item,
    update_menu_item,
    delete_menu_item,
    save_menu_recipe_line,
    delete_menu_recipe_line,
    menu_item_audit_entries,
    parse_menu_item_csv,
    validate_menu_item_import_rows,
    import_menu_items,
    get_stock_item,
    q2,
)
from settings import settings_by_section

PASSWORD_SUPPORT_MESSAGE = (
    'Need help logging in or resetting your password? '
    'Please email: pumbaskitchenapp@gmail.com'
)
DEMO_MODE = True
DEMO_COMPANY_NAME = 'Kitchen Factory Demo'
DEMO_ACCOUNT_EMAIL = 'demo@kitchenfactory.invalid'


def emergency_recovery_enabled():
    return bool(os.environ.get('MASTER_ADMIN_KEY', '').strip())


def create_app():
    app = Flask(__name__)
    app.secret_key = get_secret_key()
    app.config.update(
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE='Lax',
        SESSION_COOKIE_SECURE=(
            os.environ.get('KITCHEN_FACTORY_COOKIE_SECURE', '1' if os.environ.get('PORT') else '0')
            == '1'
        ),
        PERMANENT_SESSION_LIFETIME=__import__('datetime').timedelta(hours=12),
    )

    first_run = init_db()
    seed_data()
    app.config['FIRST_RUN'] = first_run
    demo_account_lock = Lock()

    def find_demo_user():
        demo_user_id = app.config.get('DEMO_USER_ID')
        if demo_user_id is not None:
            return get_user_by_id(demo_user_id)
        conn = get_db_connection()
        try:
            row = conn.execute(
                '''SELECT users.id FROM users
                   JOIN companies ON companies.id = users.company_id
                   WHERE companies.email = ? ORDER BY users.id LIMIT 1''',
                (DEMO_ACCOUNT_EMAIL,),
            ).fetchone()
        finally:
            conn.close()
        return get_user_by_id(row['id']) if row else None

    def ensure_demo_user():
        with demo_account_lock:
            demo_user = find_demo_user()
            if demo_user is None:
                try:
                    demo_user = register_company(
                        DEMO_COMPANY_NAME,
                        DEMO_ACCOUNT_EMAIL,
                        secrets.token_urlsafe(32),
                    )
                except ValueError:
                    demo_user = find_demo_user()
                    if demo_user is None:
                        raise
            conn = get_db_connection()
            try:
                conn.execute(
                    'UPDATE companies SET active = 1 WHERE id = ?',
                    (demo_user['company_id'],),
                )
                conn.execute(
                    '''UPDATE users SET role = 'Company Administrator', active = 1,
                              account_locked = 0, must_change_password = 0,
                              email_verified = 1 WHERE id = ?''',
                    (demo_user['id'],),
                )
                conn.commit()
            finally:
                conn.close()
            demo_user = get_user_by_id(demo_user['id'])
            if demo_user is None:
                raise RuntimeError('The demo account could not be loaded.')
            app.config['DEMO_USER_ID'] = demo_user['id']
            return demo_user

    @app.context_processor
    def inject_settings():
        return {
            'settings': get_settings(),
            'csrf_token': session.get('csrf_token', ''),
            'demo_mode': DEMO_MODE,
        }

    @app.before_request
    def require_login():
        if DEMO_MODE:
            demo_user = ensure_demo_user()
            session['user_id'] = demo_user['id']
            session['company_id'] = demo_user['company_id']
            session['username'] = demo_user['email']
            session['must_change_password'] = False
            session.permanent = True
            set_tenant_context(demo_user['company_id'], demo_user['id'])
            return None

        incoming_session_cookie = bool(
            request.cookies.get(app.config['SESSION_COOKIE_NAME'])
        )
        csrf_token_created = 'csrf_token' not in session
        if csrf_token_created:
            session['csrf_token'] = secrets.token_urlsafe(32)
        if request.method == 'POST':
            form_token = request.form.get('csrf_token', '')
            header_token = request.headers.get('X-CSRF-Token', '')
            submitted = form_token or header_token
            expected = session.get('csrf_token', '')
            failure_reason = None
            if not expected:
                failure_reason = 'session token missing'
            elif not submitted:
                failure_reason = 'submitted token missing'
            elif not hmac.compare_digest(submitted, expected):
                failure_reason = 'submitted token does not match session token'

            development_mode = os.environ.get('FLASK_ENV') == 'development'
            validation = 'failed' if failure_reason else 'succeeded'
            action = (
                'bypassed in development'
                if failure_reason and development_mode
                else 'rejected' if failure_reason
                else 'accepted'
            )
            log_method = app.logger.warning if failure_reason else app.logger.info
            log_method(
                'CSRF validation %s; action=%s reason=%s method=%s endpoint=%s '
                'incoming_session_cookie=%s csrf_token_created=%s '
                'submitted_source=%s expected_token_length=%d submitted_token_length=%d',
                validation,
                action,
                failure_reason or 'none',
                request.method,
                request.endpoint,
                incoming_session_cookie,
                csrf_token_created,
                'form' if form_token else 'header' if header_token else 'missing',
                len(expected) if isinstance(expected, str) else -1,
                len(submitted),
            )
            if failure_reason:
                if not development_mode:
                    flash('Your session security token expired. Please try again.', 'error')
                    return redirect(url_for('login'))
        elif csrf_token_created:
            app.logger.info(
                'CSRF token created; method=%s endpoint=%s incoming_session_cookie=%s '
                'csrf_token_created=True',
                request.method,
                request.endpoint,
                incoming_session_cookie,
            )
        public_routes = {
            'login', 'register', 'forgot_password', 'reset_password_view',
            'verify_email', 'resend_verification', 'static',
        }
        if request.endpoint in public_routes:
            return None
        user = get_user_by_id(session.get('user_id'))
        if (
            not user or not user['active'] or not user['company_active']
            or user['account_locked']
        ):
            session.clear()
            return redirect(url_for('login'))
        if user['must_change_password'] and request.endpoint not in {
            'change_password', 'logout', 'static',
        }:
            return redirect(url_for('change_password'))
        g_user = dict(user)
        request.environ['kitchen_factory_user'] = g_user
        set_tenant_context(user['company_id'], user['id'])
        if request.endpoint == 'users' and user['role'] != 'Company Administrator':
            flash('Only a Company Administrator can manage users.', 'error')
            return redirect(url_for('dashboard'))
        return None

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if DEMO_MODE and request.method == 'POST':
            return redirect(url_for('dashboard'))
        if request.method == 'POST':
            email = request.form.get('email', '').strip()
            user = authenticate_user(email, request.form.get('password', ''))
            if user:
                session.clear()
                session['user_id'] = user['id']
                session['company_id'] = user['company_id']
                session['username'] = user['email']
                session['must_change_password'] = bool(user['must_change_password'])
                session.permanent = True
                set_tenant_context(user['company_id'], user['id'])
                if user['must_change_password']:
                    flash('Sign in with a temporary password. Choose a new password to continue.', 'success')
                    return redirect(url_for('change_password'))
                flash('Logged in successfully.', 'success')
                return redirect(url_for('dashboard'))
            flash('Invalid email address or password.', 'error')
        return render_template('login.html')

    @app.route('/register', methods=['GET', 'POST'])
    def register():
        if DEMO_MODE:
            return redirect(url_for('dashboard'))
        if request.method == 'POST':
            try:
                user = register_company(
                    request.form.get('company_name', ''),
                    request.form.get('email', ''),
                    request.form.get('password', ''),
                )
                session.clear()
                session['user_id'] = user['id']
                session['company_id'] = user['company_id']
                session['username'] = user['email']
                session.permanent = True
                set_tenant_context(user['company_id'], user['id'])
                flash('Company registered successfully. You are now signed in.', 'success')
                return redirect(url_for('dashboard'))
            except ValueError as exc:
                flash(str(exc), 'error')
        return render_template('register.html')

    @app.route('/forgot-password', methods=['GET', 'POST'])
    def forgot_password():
        flash(PASSWORD_SUPPORT_MESSAGE, 'info')
        return redirect(url_for('login'))

    @app.route('/verify-email/<token>')
    def verify_email(token):
        del token
        flash('Email verification is no longer required. You can sign in.', 'success')
        return redirect(url_for('login'))

    @app.route('/resend-verification', methods=['GET', 'POST'])
    def resend_verification():
        flash('Email verification is no longer required. You can sign in.', 'success')
        return redirect(url_for('login'))

    @app.route('/reset-password/<token>', methods=['GET', 'POST'])
    def reset_password_view(token):
        del token
        flash(PASSWORD_SUPPORT_MESSAGE, 'info')
        return redirect(url_for('login'))

    @app.route('/change-password', methods=['GET', 'POST'])
    def change_password():
        user_id = session.get('user_id')
        if not user_id:
            return redirect(url_for('login'))
        if request.method == 'POST':
            try:
                change_temporary_password(
                    user_id, request.form.get('password', ''),
                    request.form.get('confirm_password', ''),
                )
                session['must_change_password'] = False
                flash('Password updated. You can now continue.', 'success')
                return redirect(url_for('dashboard'))
            except ValueError as exc:
                flash(str(exc), 'error')
        return render_template('change_password.html')

    @app.route('/logout', methods=['POST'])
    def logout():
        if DEMO_MODE:
            return redirect(url_for('dashboard'))
        session.clear()
        clear_tenant_context()
        return redirect(url_for('login'))

    @app.teardown_request
    def reset_request_context(error=None):
        clear_tenant_context()

    @app.route('/about')
    def about():
        return render_template('about.html')

    @app.route('/')
    @app.route('/dashboard')
    def dashboard():
        metrics = dashboard_metrics()
        recent = audit_entries(8)
        stock = list_stock_items()
        return render_template('dashboard.html', user=get_user_by_username(session['username']), metrics=metrics, recent=recent, stock=stock)

    @app.route('/stock-items', methods=['GET', 'POST'])
    def stock_items():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            try:
                item_id = create_stock_item(
                    request.form.get('code', ''),
                    request.form.get('name', ''),
                    request.form.get('unit', 'kg'),
                    request.form.get('quantity', '0'),
                    request.form.get('category_id', None),
                    user['id'],
                    unit_cost=request.form.get('unit_cost', '0'),
                    item_type=request.form.get('item_type', 'raw_material'),
                )
                add_audit(user['id'], 'STOCK ITEM CREATED', 'Stock Items', item_id, f'Created stock item {request.form.get("code")}')
                flash('Stock item created successfully.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        return render_template('stock_items.html', stock=list_stock_items(), categories=list_categories(), user=user)

    @app.route('/stock-items/<int:item_id>/status', methods=['POST'])
    def update_stock_item_status(item_id):
        user = get_user_by_username(session['username'])
        try:
            active = int(request.form.get('active', ''))
            code = set_stock_item_active(item_id, active, user['id'])
            status = 'active' if active else 'inactive'
            add_audit(user['id'], 'STOCK ITEM STATUS UPDATED', 'Stock Items', item_id, f'Set {code} to {status}')
            flash(f'Stock item {code} set to {status}.', 'success')
        except (TypeError, ValueError) as exc:
            flash(str(exc), 'error')
        return redirect(url_for('stock_items'))

    @app.route('/stock-items/edit/<int:item_id>', methods=['GET', 'POST'])
    def edit_stock_item(item_id):
        user = get_user_by_username(session['username'])
        item = get_stock_item(item_id)
        if request.method == 'POST':
            try:
                update_stock_item(
                    item_id,
                    request.form.get('code', ''),
                    request.form.get('name', ''),
                    request.form.get('unit', 'kg'),
                    request.form.get('quantity', '0'),
                    request.form.get('category_id', None),
                    user['id'],
                    active=request.form.get('active', 'on') == 'on',
                    unit_cost=request.form.get('unit_cost', '0'),
                    notes=request.form.get('notes', ''),
                    item_type=request.form.get('item_type', 'raw_material'),
                )
                add_audit(user['id'], 'STOCK ITEM EDITED', 'Stock Items', item_id, f'Updated stock item {request.form.get("code")}')
                flash('Stock item updated.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
            item = get_stock_item(item_id)
        item = dict(item)
        item['total_cost'] = q2(
            Decimal(str(item['quantity'])) * Decimal(str(item['unit_cost']))
        )
        return render_template('stock_item_edit.html', item=item, categories=list_categories(), user=user)

    @app.route('/import-stock', methods=['GET', 'POST'])
    def import_stock():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            csv_file = request.files.get('csv_file')
            if not csv_file:
                flash('Please upload a CSV file.', 'error')
                return render_template('import_stock.html', user=user)
            try:
                rows = parse_csv(csv_file.read().decode('utf-8-sig'))
                valid_rows, errors = validate_import_rows(rows)
                if errors:
                    return render_template('import_stock.html', user=user, preview=valid_rows, errors=errors)
                accepted, rejected = import_stock_items(valid_rows, user['id'])
                add_audit(user['id'], 'STOCK ITEM IMPORTED', 'Stock Items', 'import', f'Imported {len(accepted)} stock items')
                flash(f'Import complete: {len(accepted)} accepted, {len(rejected)} rejected.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        return render_template('import_stock.html', user=user)

    @app.route('/categories', methods=['GET', 'POST'])
    def categories():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            try:
                category_id = create_category(request.form.get('name', ''), request.form.get('description', ''))
                add_audit(user['id'], 'CATEGORY CREATED', 'Settings', category_id, f'Created category {request.form.get("name")}')
                flash('Category created.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        conn = get_db_connection()
        all_categories = conn.execute('SELECT * FROM categories ORDER BY name').fetchall()
        conn.close()
        return render_template('categories.html', categories=all_categories, user=user)

    @app.route('/categories/<int:category_id>/edit', methods=['POST'])
    def edit_category(category_id):
        user = get_user_by_username(session['username'])
        try:
            update_category(category_id, request.form.get('name', ''), request.form.get('description', ''))
            add_audit(user['id'], 'CATEGORY UPDATED', 'Settings', category_id, 'Updated category')
            flash('Category updated.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('categories'))

    @app.route('/categories/<int:category_id>/status', methods=['POST'])
    def category_status(category_id):
        user = get_user_by_username(session['username'])
        try:
            active = int(request.form.get('active', '0'))
            name = set_category_active(category_id, active)
            add_audit(user['id'], 'CATEGORY STATUS UPDATED', 'Settings', category_id, f'Updated {name} status')
            flash('Category status updated.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('categories'))

    @app.route('/menuitems')
    @app.route('/menu-items')
    def menu_items():
        user = get_user_by_username(session['username'])
        search = request.args.get('search', '').strip()
        return render_template(
            'menu_items.html', user=user, menu_items=list_menu_items(search), search=search
        )

    @app.route('/menu-items/export.csv')
    def export_menu_items():
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            'Code', 'Name', 'Type', 'Category', 'Cost Price', 'Selling Price',
            'Gross Profit', 'Gross Profit %', 'Active',
        ])
        for item in list_menu_items(request.args.get('search', '').strip()):
            writer.writerow([
                item['code'], item['name'], item['item_type'], item['category_name'],
                item['cost_price'], item['selling_price'], item['gross_profit'],
                item['gross_profit_percentage'], 'Yes' if item['active'] else 'No',
            ])
        response = app.response_class(output.getvalue(), mimetype='text/csv')
        response.headers['Content-Disposition'] = 'attachment; filename=menu_items.csv'
        return response

    @app.route('/menu-items/import-template.csv')
    def menu_items_import_template():
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Code', 'Name', 'Type', 'Category', 'Selling Price', 'Active'])
        response = app.response_class(output.getvalue(), mimetype='text/csv')
        response.headers['Content-Disposition'] = 'attachment; filename=menu_items_import_template.csv'
        return response

    @app.route('/menu-items/import', methods=['GET', 'POST'])
    def import_menu_items_view():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            csv_file = request.files.get('csv_file')
            if not csv_file or not csv_file.filename:
                flash('Please select a CSV file to import.', 'error')
                return render_template('import_menu_items.html', user=user)
            try:
                rows = parse_menu_item_csv(csv_file.read().decode('utf-8-sig'))
                valid_rows, errors = validate_menu_item_import_rows(rows)
                if errors:
                    return render_template(
                        'import_menu_items.html', user=user,
                        preview=valid_rows, errors=errors,
                    )
                imported_ids = import_menu_items(rows, user['id'])
                flash(f'Imported {len(imported_ids)} menu items successfully.', 'success')
                return redirect(url_for('menu_items'))
            except (UnicodeDecodeError, ValueError) as exc:
                flash(str(exc), 'error')
        return render_template('import_menu_items.html', user=user)

    @app.route('/menu-items/new', methods=['GET', 'POST'])
    def create_menu_item_view():
        user = get_user_by_username(session['username'])
        categories = list_menu_categories(include_inactive=False)
        if request.method == 'POST':
            try:
                item_id = create_menu_item(
                    request.form.get('code', ''), request.form.get('name', ''),
                    request.form.get('item_type', ''), request.form.get('category_id', ''),
                    request.form.get('selling_price', ''), user['id'],
                )
                flash('Menu item created. Add its recipe to calculate cost.', 'success')
                return redirect(url_for('menu_item_detail', menu_item_id=item_id, tab='recipe'))
            except ValueError as exc:
                flash(str(exc), 'error')
        return render_template(
            'menu_item_detail.html', user=user, item=None, categories=categories,
            recipe=[], audit_entries=[], stock_items=list_stock_items(), tab='general',
        )

    @app.route('/menu-items/<int:menu_item_id>', methods=['GET', 'POST'])
    def menu_item_detail(menu_item_id):
        user = get_user_by_username(session['username'])
        item = get_menu_item(menu_item_id)
        if item is None:
            flash('Menu item not found.', 'error')
            return redirect(url_for('menu_items'))
        if request.method == 'POST':
            try:
                update_menu_item(
                    menu_item_id, request.form.get('code', ''), request.form.get('name', ''),
                    request.form.get('item_type', ''), request.form.get('category_id', ''),
                    request.form.get('selling_price', ''), request.form.get('active') == '1',
                    user['id'],
                )
                flash('Menu item updated.', 'success')
                return redirect(url_for('menu_item_detail', menu_item_id=menu_item_id))
            except ValueError as exc:
                flash(str(exc), 'error')
        categories = list_menu_categories()
        recipe = get_menu_item_recipe(menu_item_id)
        tab = request.args.get('tab', 'general')
        if tab not in {'general', 'recipe', 'audit'}:
            tab = 'general'
        return render_template(
            'menu_item_detail.html', user=user, item=item, categories=categories,
            recipe=recipe, audit_entries=menu_item_audit_entries(menu_item_id),
            stock_items=list_stock_items(), tab=tab,
            read_only=request.args.get('mode') == 'view',
        )

    @app.route('/menu-items/<int:menu_item_id>/delete', methods=['POST'])
    def delete_menu_item_view(menu_item_id):
        user = get_user_by_username(session['username'])
        try:
            delete_menu_item(menu_item_id, user['id'])
            flash('Menu item deleted.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_items'))

    @app.route('/menu-items/<int:menu_item_id>/recipe', methods=['POST'])
    def save_menu_item_recipe(menu_item_id):
        user = get_user_by_username(session['username'])
        try:
            line_id = request.form.get('line_id', '').strip()
            save_menu_recipe_line(
                menu_item_id, request.form.get('item_type', ''),
                request.form.get('stock_item_id', ''), request.form.get('quantity', ''),
                user['id'], int(line_id) if line_id else None,
            )
            flash('Recipe line saved.', 'success')
        except (TypeError, ValueError) as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_item_detail', menu_item_id=menu_item_id, tab='recipe'))

    @app.route('/menu-items/<int:menu_item_id>/recipe/<int:line_id>/delete', methods=['POST'])
    def remove_menu_item_recipe_line(menu_item_id, line_id):
        user = get_user_by_username(session['username'])
        try:
            delete_menu_recipe_line(menu_item_id, line_id, user['id'])
            flash('Recipe line deleted.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_item_detail', menu_item_id=menu_item_id, tab='recipe'))

    @app.route('/settings/menu-categories', methods=['GET', 'POST'])
    def menu_categories():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            try:
                create_menu_category(request.form.get('name', ''), user['id'])
                flash('Menu category created.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
            return redirect(url_for('menu_categories'))
        return render_template(
            'menu_categories.html', user=user,
            categories=list_menu_categories(),
        )

    @app.route('/settings/menu-categories/<int:category_id>/edit', methods=['POST'])
    def edit_menu_category(category_id):
        user = get_user_by_username(session['username'])
        try:
            update_menu_category(category_id, request.form.get('name', ''), user['id'])
            flash('Menu category updated.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_categories'))

    @app.route('/settings/menu-categories/<int:category_id>/status', methods=['POST'])
    def menu_category_status(category_id):
        user = get_user_by_username(session['username'])
        try:
            active = int(request.form.get('active', '0'))
            set_menu_category_active(category_id, active, user['id'])
            flash('Menu category status updated.', 'success')
        except (TypeError, ValueError) as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_categories'))

    @app.route('/settings/menu-categories/<int:category_id>/delete', methods=['POST'])
    def remove_menu_category(category_id):
        try:
            delete_menu_category(category_id)
            flash('Menu category deleted.', 'success')
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('menu_categories'))

    @app.route('/manufacturing', methods=['GET', 'POST'])
    def manufacturing():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            output_item_id = request.form.get('output_item_id')
            ingredients = []
            for key, value in request.form.items():
                if key.startswith('ingredient_') and key.endswith('_qty'):
                    item_id = key.replace('ingredient_', '').replace('_qty', '')
                    qty = value.strip()
                    if qty and qty != '0':
                        ingredients.append({'item_id': int(item_id), 'quantity': qty})
            try:
                actual_output = request.form.get('actual_output')
                result = create_manufacturing_transaction(
                    output_item_id, ingredients, actual_output, actual_output, user['id'],
                    request.form.get('notes', ''), confirmation=request.form.get('confirmation') == '1'
                )
                add_audit(user['id'], 'MANUFACTURING CREATED', 'Manufacturing', result['transaction_id'], f'Manufactured output assigned to stock item id {output_item_id}')
                flash('Manufacturing transaction saved successfully.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        stock = list_stock_items()
        return render_template(
            'manufacturing.html',
            stock=stock,
            output_stock=[item for item in stock if item['item_type'] == 'manufactured_item'],
            user=user,
            allow_negative_stock=get_setting('allow_negative_stock', '0') == '1',
            history=list_manufacturing_history(10),
            calculator_mode=get_setting('calculator_mode', '0') == '1',
            require_notes=get_setting('manufacturing_require_notes', '0') == '1',
            require_confirmation=get_setting('manufacturing_require_confirmation', '0') == '1',
        )

    @app.route('/portioning', methods=['GET', 'POST'])
    def portioning():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            action = request.form.get('action')
            try:
                if action == 'bulk':
                    result = create_bulk_portioning(
                        request.form.get('source_item_id'),
                        request.form.get('destination_item_id'),
                        request.form.get('quantity_portioned'),
                        user['id'],
                        request.form.get('notes', ''), request.form.get('waste_reason', '')
                    )
                    add_audit(user['id'], 'PORTIONING CREATED', 'Portioning', result['session_id'], 'Bulk portioning created')
                    flash('Bulk portioning recorded.', 'success')
                elif action == 'yield_loss':
                    result = create_yield_loss_portioning(
                        request.form.get('source_item_id'),
                        request.form.get('destination_item_id') or None,
                        request.form.get('original_quantity'),
                        request.form.get('usable_quantity'),
                        request.form.get('original_total_cost'),
                        user['id'],
                        request.form.get('notes', ''), request.form.get('waste_reason', '')
                    )
                    add_audit(user['id'], 'PORTIONING COMPLETED', 'Portioning', result['session_id'], 'Yield loss transaction recorded')
                    message = f'Yield loss recorded. Adjusted cost: {result["adjusted_cost_per_unit"]}'
                    if result.get('yield_warning'):
                        message += ' Warning: yield is below the configured minimum.'
                    flash(message, 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        stock = list_stock_items()
        return render_template('portioning.html', stock=stock,
                               portioned_stock=[item for item in stock if item['item_type'] == 'portioned_item'],
                               recent=list_recent_portioning(10), user=user,
                               require_waste_reason=get_setting('portioning_require_waste_reason', '0') == '1')

    @app.route('/reports')
    def reports():
        user = get_user_by_username(session['username'])
        values = get_settings()
        return render_template('reports.html', user=user, manufacturing_history=list_manufacturing_history(20),
                               recent_portioning=list_recent_portioning(20), report_settings=values)

    @app.route('/reports/stock-items')
    def stock_item_report_view():
        user = get_user_by_username(session['username'])
        filters = {
            'search': request.args.get('search', '').strip(),
            'category': request.args.get('category', '').strip(),
            'item_type': request.args.get('item_type', '').strip(),
            'unit': request.args.get('unit', '').strip(),
            'status': request.args.get('status', '').strip(),
            'sort': request.args.get('sort', 'code'),
            'direction': request.args.get('direction', 'asc'),
        }
        try:
            page = max(1, int(request.args.get('page', '1')))
        except ValueError:
            page = 1
        per_page = 25
        rows, totals, total_count = stock_item_report(**filters, page=page, per_page=per_page)
        pages = max(1, (total_count + per_page - 1) // per_page)
        return render_template(
            'stock_item_report.html',
            user=user,
            rows=rows,
            totals=totals,
            total_count=total_count,
            page=page,
            pages=pages,
            filters=filters,
            categories=list_categories(),
            generated_at=datetime.now().strftime('%Y-%m-%d %H:%M'),
            report_settings=get_settings(),
        )

    @app.route('/reports/stock-items/export/<export_format>')
    def export_stock_item_report(export_format):
        user = get_user_by_username(session['username'])
        filters = {
            'search': request.args.get('search', '').strip(),
            'category': request.args.get('category', '').strip(),
            'item_type': request.args.get('item_type', '').strip(),
            'unit': request.args.get('unit', '').strip(),
            'status': request.args.get('status', '').strip(),
            'sort': request.args.get('sort', 'code'),
            'direction': request.args.get('direction', 'asc'),
        }
        rows, totals, _ = stock_item_report(**filters, page=1, per_page=1000000)
        headers = ['Stock Code', 'Stock Name', 'Stock Type', 'Unit', 'Quantity', 'Cost Per Unit', 'Total Cost', 'Category', 'Status', 'Last Updated']
        data = [[row['code'], row['name'], row['item_type'].replace('_', ' ').title(), row['unit'], row['quantity'], row['unit_cost'], row['total_cost'], row['category_name'] or '', 'Active' if row['active'] else 'Inactive', row['date_modified']] for row in rows]
        currency = get_settings().get('reporting_currency', 'R')
        filename = 'stock_item_report'
        if export_format == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(headers)
            writer.writerows(data)
            response = app.response_class(output.getvalue(), mimetype='text/csv')
            response.headers['Content-Disposition'] = f'attachment; filename={filename}.csv'
            return response
        if export_format == 'xlsx':
            from openpyxl import Workbook
            output = io.BytesIO()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Stock Item Report'
            sheet.append(['Kitchen Factory', 'Stock Item Report'])
            sheet.append(['Generated', datetime.now().strftime('%Y-%m-%d %H:%M')])
            sheet.append(['Generated By', user['display_name']])
            sheet.append([])
            sheet.append(headers)
            for row in data:
                excel_row = list(row)
                excel_row[5] = float(Decimal(str(row[5])))
                excel_row[6] = float(Decimal(str(row[6])))
                sheet.append(excel_row)
            sheet.append([])
            sheet.append(['Total Stock Items', totals['stock_items']])
            sheet.append(['Total Active Items', totals['active_items']])
            sheet.append(['Total Inactive Items', totals['inactive_items']])
            sheet.append(['Total Inventory Value', float(totals['inventory_value'])])
            for row_number in range(6, 6 + len(data)):
                sheet.cell(row=row_number, column=6).number_format = f'"{currency}"#,##0.00'
                sheet.cell(row=row_number, column=7).number_format = f'"{currency}"#,##0.00'
            sheet.cell(row=sheet.max_row, column=2).number_format = f'"{currency}"#,##0.00'
            workbook.save(output)
            response = app.response_class(output.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response.headers['Content-Disposition'] = f'attachment; filename={filename}.xlsx'
            return response
        if export_format == 'pdf':
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            output = io.BytesIO()
            document = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=0.3 * inch, leftMargin=0.3 * inch)
            styles = getSampleStyleSheet()
            body = [Paragraph('Kitchen Factory - Stock Item Report', styles['Title']), Paragraph(f'Generated: {datetime.now():%Y-%m-%d %H:%M} | Generated By: {user["display_name"]}', styles['Normal']), Spacer(1, 0.15 * inch)]
            pdf_data = []
            for row in data:
                formatted = [str(value) for value in row]
                formatted[5] = f'{currency}{Decimal(str(row[5])):,.2f}'
                formatted[6] = f'{currency}{Decimal(str(row[6])):,.2f}'
                pdf_data.append(formatted)
            pdf_rows = [headers] + pdf_data
            table = Table(pdf_rows, repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#132238')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('GRID', (0, 0), (-1, -1), 0.25, colors.grey), ('FONTSIZE', (0, 0), (-1, -1), 7), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f6f8')])]))
            body.append(table)
            body.append(Spacer(1, 0.15 * inch))
            body.append(Paragraph(f'Total Stock Items: {totals["stock_items"]} | Active: {totals["active_items"]} | Inactive: {totals["inactive_items"]} | Inventory Value: {currency}{totals["inventory_value"]}', styles['Normal']))
            document.build(body)
            response = app.response_class(output.getvalue(), mimetype='application/pdf')
            response.headers['Content-Disposition'] = f'attachment; filename={filename}.pdf'
            return response
        flash('Unsupported stock item report export format.', 'error')
        return redirect(url_for('stock_item_report_view'))

    @app.route('/audit-log')
    def audit_log():
        user = get_user_by_username(session['username'])
        return render_template('audit_log.html', user=user, entries=audit_entries(50))

    @app.route('/users', methods=['GET', 'POST'])
    def users():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            try:
                user_id = request.form.get('user_id')
                if user_id:
                    update_user(int(user_id), request.form.get('email', ''), request.form.get('display_name', ''),
                                request.form.get('role', ''), request.form.get('active') == '1')
                    action = 'USER UPDATED'
                else:
                    user_id = create_user(
                        request.form.get('email', ''), request.form.get('display_name', ''),
                        request.form.get('role', ''), request.form.get('password', ''),
                        request.form.get('confirm_password', ''),
                        request.form.get('active') == '1',
                    )
                    action = 'USER CREATED'
                add_audit(user['id'], action, 'Users', user_id, 'User administration change')
                flash('User saved.', 'success')
            except (TypeError, ValueError) as exc:
                flash(str(exc), 'error')
            return redirect(url_for('users'))
        return render_template(
            'users.html', user=user, users=list_users(),
            recovery_enabled=emergency_recovery_enabled(),
        )

    @app.route('/admin-recovery', methods=['GET', 'POST'])
    def admin_recovery():
        if not emergency_recovery_enabled():
            return app.make_response(('Not found', 404))
        administrator = get_user_by_username(session['username'])
        if administrator['role'] != 'Company Administrator':
            flash('Only a Company Administrator can access recovery.', 'error')
            return redirect(url_for('dashboard'))
        recovery_active = session.get('admin_recovery_until', 0) > time.time()
        if request.method == 'POST':
            action = request.form.get('action', '')
            try:
                try:
                    target_user_id = int(request.form.get('target_user_id', ''))
                except (TypeError, ValueError):
                    try:
                        emergency_admin_recovery(
                            administrator['id'], 0, 'invalid_access_attempt',
                            request.form.get('administrator_password', ''),
                            request.form.get('master_admin_key', ''),
                        )
                    except ValueError:
                        pass
                    raise ValueError('Select a valid target user.')
                if action != 'authorize_access' and not recovery_active:
                    try:
                        emergency_admin_recovery(
                            administrator['id'], target_user_id,
                            'invalid_access_attempt',
                            request.form.get('administrator_password', ''),
                            request.form.get('master_admin_key', ''),
                        )
                    except ValueError:
                        pass
                    raise ValueError(
                        'Verify your administrator password and recovery key to open recovery.'
                    )
                target_email = emergency_admin_recovery(
                    administrator['id'],
                    target_user_id,
                    action,
                    request.form.get('administrator_password', ''),
                    request.form.get('master_admin_key', ''),
                    request.form.get('new_password'),
                    request.form.get('confirm_password'),
                    request.form.get('force_password_change') == '1',
                )
                if action == 'authorize_access':
                    session['admin_recovery_until'] = time.time() + 300
                    flash(
                        'Emergency recovery access granted for five minutes.',
                        'success',
                    )
                else:
                    flash(
                        f'Recovery action completed for {target_email}. '
                        'The action was recorded in the audit log.',
                        'success',
                    )
            except (TypeError, ValueError) as exc:
                flash(str(exc), 'error')
            return redirect(url_for('admin_recovery'))
        if not recovery_active:
            return render_template(
                'admin_recovery_auth.html', user=administrator,
            )
        conn = get_db_connection()
        try:
            recovery_users = conn.execute(
                '''SELECT id, email, display_name, role, active, account_locked,
                          must_change_password
                   FROM users WHERE company_id = ? ORDER BY display_name''',
                (administrator['company_id'],),
            ).fetchall()
        finally:
            conn.close()
        return render_template(
            'admin_recovery.html', user=administrator, users=recovery_users,
        )

    @app.route('/users/<int:user_id>/reset-password', methods=['POST'])
    def reset_user_password(user_id):
        user = get_user_by_username(session['username'])
        if user['role'] != 'Company Administrator':
            flash('Only a Company Administrator can reset user passwords.', 'error')
            return redirect(url_for('dashboard'))
        try:
            temporary_password = administrator_reset_user_password(user_id)
            flash(
                f'Temporary password: {temporary_password}. '
                'Share it securely; it will only be shown once.',
                'warning',
            )
        except ValueError as exc:
            flash(str(exc), 'error')
        return redirect(url_for('users'))

    @app.route('/settings', methods=['GET', 'POST'])
    def settings():
        user = get_user_by_username(session['username'])
        if request.method == 'POST':
            values = get_settings()
            for key in values:
                if key in ('allow_negative_stock', 'calculator_mode', 'manufacturing_require_notes',
                           'manufacturing_require_confirmation', 'manufacturing_auto_cost',
                           'portioning_require_waste_reason', 'audit_enabled'):
                    values[key] = '1' if request.form.get(key) == '1' else '0'
                elif key in request.form:
                    values[key] = request.form.get(key)
            try:
                update_settings(values)
                add_audit(user['id'], 'SETTING UPDATED', 'Settings', 'all', 'Updated application settings')
                flash('Settings updated.', 'success')
            except ValueError as exc:
                flash(str(exc), 'error')
        values = get_settings()
        return render_template(
            'settings.html',
            user=user,
            categories=list_categories(),
            allow_negative_stock=values.get('allow_negative_stock') == '1',
            settings=values,
            setting_sections=settings_by_section(values),
        )

    @app.route('/export-csv')
    def export_csv():
        output = io.StringIO()
        writer = csv.writer(output)
        rows = list_stock_items()
        writer.writerow(['Code', 'Name', 'Unit', 'Quantity', 'Category'])
        for row in rows:
            writer.writerow([row['code'], row['name'], row['unit'], row['quantity'], row['category_name']])
        response = app.response_class(output.getvalue(), mimetype='text/csv')
        response.headers['Content-Disposition'] = 'attachment; filename=stock_report.csv'
        return response

    return app


import os

app = create_app()

if __name__ == '__main__':
    app.run(
        debug=False,
        host='0.0.0.0',
        port=int(os.environ.get('PORT', 5000))
    )